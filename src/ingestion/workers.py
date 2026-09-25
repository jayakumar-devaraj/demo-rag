"""Polymorphic ingestion workers: every format resolves to per-page/per-tab Markdown units."""

import io
from abc import ABC, abstractmethod
from pathlib import PurePath

import docx
import openpyxl
import pymupdf
from docx.table import Table as DocxTable
from openpyxl.utils import get_column_letter
from pydantic import BaseModel


class Unit(BaseModel):
    """One hashable unit: a PDF page, a spreadsheet tab, or a whole text document."""

    page_key: str
    markdown: str


class UnsupportedFormatError(ValueError):
    """Raised for file extensions no worker handles (a caller error → 4xx)."""


def _md_cell(value: object) -> str:
    """Render one cell for a Markdown table (pipes escaped, newlines flattened)."""
    return "" if value is None else str(value).replace("|", "\\|").replace("\n", " ").strip()


def md_table(rows: list[list[object]]) -> str:
    """Render rows as a Markdown table; the first row is the header."""
    width = max(len(r) for r in rows)
    padded = [[_md_cell(c) for c in r] + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(padded[0]) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in padded[1:]]
    return "\n".join(lines)


class Worker(ABC):
    """Abstract ingestion interface: raw bytes in, ordered Markdown units out."""

    @abstractmethod
    def extract(self, data: bytes) -> list[Unit]:
        """Convert a file's bytes to Markdown units."""


class PdfWorker(Worker):
    """PDF pages → Markdown; detected table grids become Markdown tables in reading order."""

    def extract(self, data: bytes) -> list[Unit]:
        """Return one unit per page, tables and text blocks ordered top to bottom."""
        units = []
        with pymupdf.open(stream=data, filetype="pdf") as doc:
            for number, page in enumerate(doc, start=1):
                tables = page.find_tables().tables
                boxes = [pymupdf.Rect(t.bbox) for t in tables]
                parts = [(t.bbox[1], t.to_markdown().strip()) for t in tables]
                for x0, y0, x1, y1, text, *_ in page.get_text("blocks"):
                    if text.strip() and not any(pymupdf.Rect(x0, y0, x1, y1).intersects(b) for b in boxes):
                        parts.append((y0, text.strip()))
                parts.sort(key=lambda p: p[0])
                units.append(Unit(page_key=f"page:{number}", markdown="\n\n".join(p[1] for p in parts)))
        return units


class XlsxWorker(Worker):
    """Each tab → a Markdown matrix keeping Excel coordinates (row numbers, column letters)."""

    def extract(self, data: bytes) -> list[Unit]:
        """Return one unit per non-empty sheet; formulas without cached values are kept as text."""
        values = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        formulas = openpyxl.load_workbook(io.BytesIO(data), data_only=False, read_only=True)
        units = []
        for ws in values.worksheets:
            fws = formulas[ws.title]
            grid = [
                [v if v is not None else f for v, f in zip(vrow, frow, strict=True)]
                for vrow, frow in zip(ws.iter_rows(values_only=True), fws.iter_rows(values_only=True), strict=True)
            ]
            numbered = [(i, r) for i, r in enumerate(grid, start=1) if any(c is not None for c in r)]
            if not numbered:
                continue
            width = max(len(r) for _, r in numbered)
            header_idx, header = numbered[0]
            head = ["row"] + [f"{get_column_letter(c + 1)}:{_md_cell(h)}" for c, h in enumerate(header)]
            rows = [head] + [[i, *r] for i, r in numbered[1:]]
            body = md_table([r + [None] * (width + 1 - len(r)) for r in rows])
            units.append(Unit(page_key=f"tab:{ws.title}", markdown=f"## Sheet: {ws.title}\n\n{body}"))
        values.close()
        formulas.close()
        return units


class DocxWorker(Worker):
    """Word document → hierarchical Markdown: headings, paragraphs and inline tables in order."""

    def extract(self, data: bytes) -> list[Unit]:
        """Return the whole document as one unit (DOCX has no stable page boundaries)."""
        parts = []
        for block in docx.Document(io.BytesIO(data)).iter_inner_content():
            if isinstance(block, DocxTable):
                parts.append(md_table([[c.text for c in row.cells] for row in block.rows]))
            elif block.text.strip():
                style = block.style.name if block.style is not None else ""
                level = int(style.split()[-1]) if style.startswith("Heading") and style.split()[-1].isdigit() else 0
                parts.append(("#" * level + " " if level else "") + block.text.strip())
        return [Unit(page_key="page:1", markdown="\n\n".join(parts))]


class TextWorker(Worker):
    """Plain text → one unit (agreed cut: shares the DOCX single-unit text path)."""

    def extract(self, data: bytes) -> list[Unit]:
        """Decode UTF-8 (BOM tolerated) and return it as one unit."""
        return [Unit(page_key="page:1", markdown=data.decode("utf-8-sig").strip())]


WORKERS: dict[str, Worker] = {".pdf": PdfWorker(), ".xlsx": XlsxWorker(), ".docx": DocxWorker(), ".txt": TextWorker()}


def worker_for(filename: str) -> Worker:
    """Pick the worker by file extension."""
    ext = PurePath(filename).suffix.lower()
    if ext not in WORKERS:
        raise UnsupportedFormatError(f"unsupported file type: {ext or '(none)'}")
    return WORKERS[ext]
