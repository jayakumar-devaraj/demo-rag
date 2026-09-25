"""Shared test helpers: in-memory document builders and a labeled test-double embedder."""

import hashlib
import io

import docx
import openpyxl
import pymupdf

from src.db.connection import EMBEDDING_DIM


class FakeEmbedder:
    """TEST DOUBLE: deterministic hash-seeded vectors, no model call (no provider credentials)."""

    def __init__(self) -> None:
        self.calls: list[str] = []

    async def embed(self, texts: list[str]) -> list[list[float]]:
        self.calls.extend(texts)
        return [[b / 255 for b in hashlib.sha256(t.encode()).digest()] * (EMBEDDING_DIM // 32) for t in texts]


def make_pdf(pages: list[tuple[str, list[list[str]]]]) -> bytes:
    """Build a PDF: each page has a title line and a ruled table grid."""
    doc = pymupdf.open()
    for title, rows in pages:
        page = doc.new_page()
        page.insert_text((72, 60), title, fontsize=12)
        for r, row in enumerate(rows):
            for c, cell in enumerate(row):
                rect = pymupdf.Rect(72 + c * 120, 100 + r * 24, 192 + c * 120, 124 + r * 24)
                page.draw_rect(rect, color=(0, 0, 0), width=0.8)
                page.insert_text((rect.x0 + 4, rect.y0 + 16), cell, fontsize=10)
    data = doc.tobytes()
    doc.close()
    return data


def make_xlsx(sheets: dict[str, list[list[object]]]) -> bytes:
    """Build a multi-tab workbook."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_docx() -> bytes:
    """Build a DOCX with headings, paragraphs and an inline table."""
    d = docx.Document()
    d.add_heading("Quarterly Review", level=1)
    d.add_paragraph("Revenue grew in the northern region.")
    d.add_heading("Budget", level=2)
    t = d.add_table(rows=2, cols=2)
    for i, (a, b) in enumerate([("Item", "Cost"), ("Servers", "4200")]):
        t.cell(i, 0).text, t.cell(i, 1).text = a, b
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
